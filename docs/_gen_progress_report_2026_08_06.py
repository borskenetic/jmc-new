"""Generate Area 51 ITS Project Progress Report for 2026-08-06."""

from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(OUT_DIR, "Project_Progress_Report_2026-08-06.docx")


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="4" w:space="0" w:color="666666"/>'
        '  <w:left w:val="single" w:sz="4" w:space="0" w:color="666666"/>'
        '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="666666"/>'
        '  <w:right w:val="single" w:sz="4" w:space="0" w:color="666666"/>'
        "</w:tcBorders>"
    )
    tcPr.append(borders)


def set_cell_text(cell, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(2)
    set_run_font(
        p1.add_run("Area 51 Information Technology Services"),
        size=16,
        bold=True,
        color=(0x1F, 0x49, 0x7D),
    )

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    set_run_font(
        p2.add_run("Project Progress Report"),
        size=13,
        bold=True,
        color=(0x2E, 0x75, 0xB6),
    )

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(12)
    set_run_font(
        p3.add_run("JMC Attendance System  |  06 August 2026"),
        size=10,
        color=(0x59, 0x59, 0x59),
    )

    tasks = [
        {
            "no": "1",
            "module": "SMS Delivery (Multi-number)",
            "desc": (
                "Improved SMS sending to support multiple Philippine mobile numbers "
                "per student field; updated the default campus scan message template; "
                "aligned scan/attendance SMS handling for clearer entry notifications."
            ),
            "status": "Completed",
        },
        {
            "no": "2",
            "module": "SMS Blast Controls",
            "desc": (
                "Enhanced SMS Blast with filters by year, course, and section; "
                "added option to send either to student mobile numbers or emergency "
                "contacts; recipient counting and logging updated accordingly."
            ),
            "status": "Completed",
        },
        {
            "no": "3",
            "module": "Register Student (Signature)",
            "desc": (
                "Fixed registration failure (SQLSTATE 42S22 – unknown column "
                "student_signature). Diagnosed missing database column and added "
                "migration to add nullable student_signature to students and "
                "pending_students tables so signature capture on register works."
            ),
            "status": "Completed",
        },
    ]

    table = doc.add_table(rows=1 + len(tasks), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Inches(0.5), Inches(1.9), Inches(3.6), Inches(1.1)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]

    headers = ["No.", "Module / Feature", "Description / Accomplishment", "Status"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(
            cell,
            h,
            bold=True,
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=(0xFF, 0xFF, 0xFF),
        )
        shade_cell(cell, "1F497D")
        set_cell_borders(cell)

    for t in tasks:
        row = table.rows[int(t["no"])]
        alt = "F2F2F2" if int(t["no"]) % 2 == 0 else "FFFFFF"

        set_cell_text(row.cells[0], t["no"], bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], t["module"], bold=True, size=10)
        set_cell_text(row.cells[2], t["desc"], size=10)

        status_cell = row.cells[3]
        status_cell.text = ""
        sp = status_cell.paragraphs[0]
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after = Pt(4)
        set_run_font(sp.add_run("✅ Completed"), size=10, bold=True, color=(0x00, 0x70, 0x00))
        shade_cell(status_cell, "E2EFDA")

        for c in row.cells[:3]:
            shade_cell(c, alt)
            set_cell_borders(c)
        set_cell_borders(status_cell)

    doc.add_paragraph()

    notes_title = doc.add_paragraph()
    notes_title.paragraph_format.space_before = Pt(8)
    set_run_font(
        notes_title.add_run("Notes"),
        size=10,
        bold=True,
        color=(0x1F, 0x49, 0x7D),
    )

    for text in [
        "Register Student fix requires: php artisan migrate "
        "(migration 2026_08_06_000031_add_student_signature_to_students_tables).",
        "SMS modem URL (SMS_MODEM_URL) must remain configured for live SMS delivery.",
        "Project: jmc-new (JMC / pantas.org attendance system).",
    ]:
        bp = doc.add_paragraph(style="List Bullet")
        set_run_font(bp.add_run(text), size=9, color=(0x40, 0x40, 0x40))

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(16)
    set_run_font(
        footer.add_run("Prepared for daily progress tracking — Area 51 ITS"),
        size=8,
        color=(0x80, 0x80, 0x80),
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
